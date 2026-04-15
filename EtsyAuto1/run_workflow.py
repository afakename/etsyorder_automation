"""
run_workflow.py
Unified workflow runner — fetches orders from both Etsy and Shopify,
processes them through FilenameGenerator + FileDatabase, and writes
a combined Excel sheet with Source column and per-platform color coding.

Usage:
    python run_workflow.py            # last 7 days
    python run_workflow.py --days 30  # last 30 days
"""
import sys
import argparse
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()

from config import Config
from logger import setup_logger
from filename_generator import FilenameGenerator
from file_database import FileDatabase
from connectors.etsy import EtsyConnector
from connectors.shopify import ShopifyConnector

# openpyxl for formatting
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Colour palette ────────────────────────────────────────────────────────────
ETSY_FILL   = PatternFill("solid", fgColor="F56400")   # Etsy orange
SHOPIFY_FILL = PatternFill("solid", fgColor="96BF48")  # Shopify green
WHITE_FONT  = Font(color="FFFFFF", bold=True, size=9)
HEADER_FILL = PatternFill("solid", fgColor="2D2D2D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL    = PatternFill("solid", fgColor="F7F7F7")
THIN_BORDER = Border(
    bottom=Side(style='thin', color='DDDDDD')
)

COLUMNS = [
    "Source", "Order #", "Order Date", "Customer",
    "SKU", "Type", "Personalization", "Center Design", "Year",
    "Qty", "Price", "Filename Generated", "File Found", "File Path",
    "Message from Buyer",
]


def setup():
    Config.get_log_path().mkdir(parents=True, exist_ok=True)
    log_file = Config.get_log_path() / f"workflow_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
    logger = setup_logger("run_workflow", log_file)
    return logger


def fetch_all_orders(days_back: int, logger):
    """Fetch normalized orders from both platforms."""
    orders = []

    # Shopify
    shopify = ShopifyConnector()
    if shopify.is_configured():
        shopify_orders = shopify.get_recent_orders(days_back=days_back)
        logger.info(f"Shopify: {len(shopify_orders)} workflow orders")
        orders.extend(shopify_orders)
    else:
        logger.warning("Shopify not configured — skipping")

    # Etsy
    try:
        etsy = EtsyConnector()
        if etsy.is_authenticated():
            etsy_orders = etsy.get_recent_orders(days_back=days_back)
            logger.info(f"Etsy: {len(etsy_orders)} workflow orders")
            orders.extend(etsy_orders)
        else:
            logger.warning("Etsy not authenticated — skipping")
    except Exception as e:
        logger.warning(f"Etsy fetch failed: {e}")

    # Sort by order date, oldest first (queue order)
    orders.sort(key=lambda o: o.order_date or datetime.min)
    return orders


def process_orders(orders, logger):
    """Run each NormalizedItem through FilenameGenerator + FileDatabase."""
    fg = FilenameGenerator(logger)
    fd = FileDatabase(logger)

    rows = []
    for order in orders:
        for item in order.items:
            filename = fg.generate_filename_from_item(item)
            file_path = fd.find_file(filename) if filename else None

            sku_info = fg.sku_mapping.get(item.sku, {})
            product_type = sku_info.get("type", "")

            rows.append({
                "Source":              order.source.capitalize(),
                "Order #":             order.receipt_id,
                "Order Date":          order.order_date.strftime("%Y-%m-%d %H:%M") if order.order_date else "",
                "Customer":            order.customer_name,
                "SKU":                 item.sku,
                "Type":                product_type,
                "Personalization":     item.variations.get("Personalization", ""),
                "Center Design":       item.variations.get("Choose the Center Piece", ""),
                "Year":                item.variations.get("Year or No Year", ""),
                "Qty":                 item.quantity,
                "Price":               f"${item.price_usd:.2f}",
                "Filename Generated":  filename or "— could not generate —",
                "File Found":          "YES" if file_path else "NO",
                "File Path":           str(file_path) if file_path else "NOT FOUND",
                "Message from Buyer":  order.message_from_buyer or "",
            })

    logger.info(f"Processed {len(rows)} line items across {len(orders)} orders")
    return rows


def write_excel(rows, logger):
    """Write the combined worksheet with source colour coding."""
    Config.get_output_path().mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = Config.get_output_path() / f"All_Workflow_Orders_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Workflow Orders"

    # ── Header row ────────────────────────────────────────────────────────────
    ws.append(COLUMNS)
    for col_idx, _ in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 18

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, row in enumerate(rows, start=2):
        source = row["Source"]
        is_alt = (row_idx % 2 == 0)

        for col_idx, col_name in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[col_name])
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col_name == "Message from Buyer"))

            # Source column: coloured badge
            if col_name == "Source":
                cell.fill = ETSY_FILL if source == "Etsy" else SHOPIFY_FILL
                cell.font = WHITE_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name == "File Found":
                cell.font = Font(
                    bold=True,
                    color="276221" if row[col_name] == "YES" else "CC0000"
                )
            elif is_alt:
                cell.fill = ALT_FILL

    # ── Column widths ─────────────────────────────────────────────────────────
    col_widths = {
        "Source": 10, "Order #": 16, "Order Date": 18, "Customer": 22,
        "SKU": 24, "Type": 7, "Personalization": 18, "Center Design": 22,
        "Year": 14, "Qty": 5, "Price": 9,
        "Filename Generated": 30, "File Found": 11, "File Path": 50,
        "Message from Buyer": 45,
    }
    for col_idx, col_name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 15)

    ws.freeze_panes = "A2"

    wb.save(out_path)
    logger.info(f"Saved: {out_path}")
    print(f"\nSaved: {out_path}")
    return out_path


def main(days_back: int = 7):
    logger = setup()
    logger.info(f"=== Unified Workflow Run — last {days_back} days ===")

    orders = fetch_all_orders(days_back, logger)
    if not orders:
        print("No workflow orders found.")
        return

    rows = process_orders(orders, logger)
    out_path = write_excel(rows, logger)

    # Summary
    etsy_rows    = [r for r in rows if r["Source"] == "Etsy"]
    shopify_rows = [r for r in rows if r["Source"] == "Shopify"]
    found        = [r for r in rows if r["File Found"] == "YES"]
    missing      = [r for r in rows if r["File Found"] == "NO"]

    print(f"\n{'='*50}")
    print(f"  Etsy items:     {len(etsy_rows)}")
    print(f"  Shopify items:  {len(shopify_rows)}")
    print(f"  Files found:    {len(found)}")
    print(f"  Files MISSING:  {len(missing)}")
    if missing:
        print("\n  Items needing new files:")
        for r in missing:
            print(f"    [{r['Source']}] {r['Customer']} — {r['Filename Generated']}")
    print(f"{'='*50}\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--days", type=int, default=7, help="Days back to fetch orders")
    args = parser.parse_args()
    main(days_back=args.days)
